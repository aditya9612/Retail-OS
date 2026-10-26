import io
import logging
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional, Tuple

import openpyxl
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

from app.services.export_theme import get_theme_for_report


class ExcelExportService:
    """
    Reusable, production-grade Excel (.xlsx) export engine using openpyxl.

    Provides:
    - Clean workbook and worksheet initialization
    - Standardized corporate header styling (dark navy blue #1F497D, white bold text, row height)
    - Auto column width calculation based on maximum cell content length + safety margin
    - Thin borders and clean typography
    - Number/decimal/date cell formatting preservation
    - In-memory BytesIO generation and FastAPI StreamingResponse integration
    """

    # Corporate styling constants
    HEADER_FILL_COLOR = "1F497D"      # Dark navy blue
    HEADER_FONT_COLOR = "FFFFFF"      # White
    ZEBRA_FILL_COLOR = "F9FAFB"       # Very light gray / slate
    BORDER_COLOR = "D1D5DB"           # Soft gray border

    def __init__(self):
        self.header_font = Font(name="Calibri", size=11, bold=True, color=self.HEADER_FONT_COLOR)
        self.header_fill = PatternFill(
            start_color=self.HEADER_FILL_COLOR,
            end_color=self.HEADER_FILL_COLOR,
            fill_type="solid",
        )
        self.header_alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

        self.data_font = Font(name="Calibri", size=10, bold=False, color="111827")
        self.zebra_fill = PatternFill(
            start_color=self.ZEBRA_FILL_COLOR,
            end_color=self.ZEBRA_FILL_COLOR,
            fill_type="solid",
        )

        thin_side = Side(style="thin", color=self.BORDER_COLOR)
        self.cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    def create_workbook(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        enable_zebra: bool = True,
        min_col_width: int = 12,
        max_col_width: int = 50,
        report_type: str | None = None,
    ) -> openpyxl.Workbook:
        """
        Creates and styles an openpyxl Workbook with the given headers and data rows.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        # Excel sheet titles cannot exceed 31 chars
        ws.title = (sheet_title or "Export")[:31].replace(":", "").replace("/", "_").replace("\\", "_")

        theme = get_theme_for_report(report_type or sheet_title)
        header_fill = (
            PatternFill(start_color=theme.excel_header_hex, end_color=theme.excel_header_hex, fill_type="solid")
            if theme else self.header_fill
        )

        # Freeze the top header row so it stays visible during scrolling
        ws.freeze_panes = "A2"

        # 1. Header row
        ws.append(headers)
        ws.row_dimensions[1].height = 26

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = self.header_font
            cell.fill = header_fill
            cell.alignment = self.header_alignment
            cell.border = self.cell_border

        # 2. Data rows
        for row_idx, r in enumerate(rows, start=2):
            ws.append(self._sanitize_row(r))
            ws.row_dimensions[row_idx].height = 20

            is_even = (row_idx % 2 == 0)
            fill_to_use = self.zebra_fill if (enable_zebra and not is_even) else None

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = self.data_font
                cell.border = self.cell_border
                if fill_to_use:
                    cell.fill = fill_to_use

                # Format specific data types
                val = cell.value
                if isinstance(val, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (float, Decimal)):
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"
                elif isinstance(val, (datetime, date)):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = "YYYY-MM-DD"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 3. Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            col_letter = get_column_letter(col[0].column)
            calc_width = max(max_len + 4, min_col_width)
            ws.column_dimensions[col_letter].width = min(calc_width, max_col_width)

        return wb

    def _sanitize_row(self, row: list[Any]) -> list[Any]:
        """Convert objects or unsupported types to Excel-friendly representations."""
        sanitized: list[Any] = []
        for val in row:
            if val is None:
                sanitized.append("")
            elif isinstance(val, (Decimal,)):
                sanitized.append(float(val))
            elif isinstance(val, (int, float, str, bool, datetime, date)):
                sanitized.append(val)
            else:
                sanitized.append(str(val))
        return sanitized

    def export_to_stream(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        enable_zebra: bool = True,
        report_type: str | None = None,
    ) -> io.BytesIO:
        """Generates an in-memory BytesIO stream containing the complete .xlsx file."""
        wb = self.create_workbook(
            sheet_title=sheet_title,
            headers=headers,
            rows=rows,
            enable_zebra=enable_zebra,
            report_type=report_type,
        )
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    def export_to_bytes(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        enable_zebra: bool = True,
        report_type: str | None = None,
    ) -> bytes:
        """Generates raw .xlsx bytes."""
        stream = self.export_to_stream(sheet_title, headers, rows, enable_zebra, report_type=report_type)
        return stream.getvalue()

    def create_streaming_response(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        filename: str,
        enable_zebra: bool = True,
        report_type: str | None = None,
    ) -> StreamingResponse:
        """Creates a FastAPI StreamingResponse with proper Excel MIME type and Content-Disposition."""
        clean_name = filename if filename.endswith(".xlsx") else f"{filename}.xlsx"
        stream = self.export_to_stream(sheet_title, headers, rows, enable_zebra, report_type=report_type)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={clean_name}"},
        )

    def create_report_workbook(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        metadata: dict[str, Any] | None = None,
        kpis: list[dict[str, Any]] | None = None,
        branding: Any = None,
        enable_zebra: bool = True,
        report_type: str | None = None,
        presentation: Any = None,
    ) -> openpyxl.Workbook:
        """
        Creates a professional corporate Excel workbook:
        If presentation data model is provided, generates the reference-matched multi-sheet workbook.
        Otherwise falls back to the standard 2-sheet corporate workbook.
        """
        if presentation is not None:
            return self.create_reference_matched_report_workbook(presentation)

        wb = openpyxl.Workbook()

        theme = get_theme_for_report(report_type or sheet_title)
        header_hex = theme.excel_header_hex
        theme_header_fill = PatternFill(
            start_color=header_hex,
            end_color=header_hex,
            fill_type="solid",
        )

        # ----------------------------------------------------
        # SHEET 1: OVERVIEW / SUMMARY
        # ----------------------------------------------------
        ws_summary = wb.active
        ws_summary.title = "Overview"

        title_font = Font(name="Calibri", size=15, bold=True, color=header_hex)
        subtitle_font = Font(name="Calibri", size=11, bold=True, color="374151")
        section_font = Font(name="Calibri", size=11, bold=True, color=self.HEADER_FONT_COLOR)
        key_font = Font(name="Calibri", size=10, bold=True, color="1F2937")
        val_font = Font(name="Calibri", size=10, bold=False, color="111827")

        current_row = 1
        # Title
        display_title = (sheet_title or "Business Report").strip()
        ws_summary.cell(row=current_row, column=1, value=display_title).font = title_font
        current_row += 1

        # Company branding
        biz_name = getattr(branding, "business_name", None) if branding else None
        if biz_name:
            ws_summary.cell(row=current_row, column=1, value=str(biz_name)).font = subtitle_font
            current_row += 1
        biz_addr = getattr(branding, "address", None) if branding else None
        if biz_addr:
            ws_summary.cell(row=current_row, column=1, value=str(biz_addr)).font = Font(name="Calibri", size=9, italic=True, color="555555")
            current_row += 1
        _biz_contact_parts = []
        if branding:
            if getattr(branding, "phone", None):
                _biz_contact_parts.append(f"Tel: {branding.phone}")
            if getattr(branding, "email", None):
                _biz_contact_parts.append(f"Email: {branding.email}")
            if getattr(branding, "show_gstin", True) and getattr(branding, "gstin", None):
                _biz_contact_parts.append(f"GSTIN: {branding.gstin}")
        if _biz_contact_parts:
            ws_summary.cell(row=current_row, column=1, value=" | ".join(_biz_contact_parts)).font = Font(name="Calibri", size=9, color="666666")
            current_row += 1

        current_row += 1  # blank line

        # Metadata Section
        if metadata:
            ws_summary.cell(row=current_row, column=1, value="REPORT PARAMETERS").font = section_font
            ws_summary.cell(row=current_row, column=1).fill = theme_header_fill
            ws_summary.cell(row=current_row, column=2, value="DETAILS").font = section_font
            ws_summary.cell(row=current_row, column=2).fill = theme_header_fill
            ws_summary.row_dimensions[current_row].height = 22
            current_row += 1

            for k, v in metadata.items():
                c1 = ws_summary.cell(row=current_row, column=1, value=str(k))
                c2 = ws_summary.cell(row=current_row, column=2, value=str(v if v is not None else "N/A"))
                c1.font = key_font
                c2.font = val_font
                c1.border = self.cell_border
                c2.border = self.cell_border
                current_row += 1

            current_row += 1  # blank line

        # KPI Section
        if kpis:
            ws_summary.cell(row=current_row, column=1, value="KEY METRICS (KPIs)").font = section_font
            ws_summary.cell(row=current_row, column=1).fill = theme_header_fill
            ws_summary.cell(row=current_row, column=2, value="VALUE").font = section_font
            ws_summary.cell(row=current_row, column=2).fill = theme_header_fill
            ws_summary.row_dimensions[current_row].height = 22
            current_row += 1

            for item in kpis:
                lbl = item.get("label") or item.get("metric") or "Metric"
                val = item.get("value")
                c1 = ws_summary.cell(row=current_row, column=1, value=str(lbl))
                c2 = ws_summary.cell(row=current_row, column=2, value=val)
                c1.font = key_font
                c2.font = val_font
                c1.border = self.cell_border
                c2.border = self.cell_border

                if isinstance(val, (int, float, Decimal)):
                    c2.alignment = Alignment(horizontal="right", vertical="center")
                    c2.number_format = "#,##0.00" if isinstance(val, (float, Decimal)) else "#,##0"
                else:
                    c2.alignment = Alignment(horizontal="left", vertical="center")
                current_row += 1

        # Auto-adjust summary column widths
        for col in ws_summary.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            col_letter = get_column_letter(col[0].column)
            ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 22)

        # ----------------------------------------------------
        # SHEET 2: DETAILED DATA
        # ----------------------------------------------------
        data_sheet_title = (sheet_title or "Details")[:31].replace(":", "").replace("/", "_").replace("\\", "_")
        if data_sheet_title.lower() == "overview":
            data_sheet_title = "Data"
        ws_data = wb.create_sheet(title=data_sheet_title)
        ws_data.freeze_panes = "A2"

        # Headers
        ws_data.append(headers)
        ws_data.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.font = self.header_font
            cell.fill = theme_header_fill
            cell.alignment = self.header_alignment
            cell.border = self.cell_border

        # Data rows
        for row_idx, r in enumerate(rows, start=2):
            ws_data.append(self._sanitize_row(r))
            ws_data.row_dimensions[row_idx].height = 20

            is_even = (row_idx % 2 == 0)
            fill_to_use = self.zebra_fill if (enable_zebra and not is_even) else None

            for col_idx in range(1, len(headers) + 1):
                cell = ws_data.cell(row=row_idx, column=col_idx)
                cell.font = self.data_font
                cell.border = self.cell_border
                if fill_to_use:
                    cell.fill = fill_to_use

                val = cell.value
                if isinstance(val, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (float, Decimal)):
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"
                elif isinstance(val, (datetime, date)):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = "YYYY-MM-DD"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-filter & column widths
        if len(headers) > 0 and len(rows) > 0:
            ws_data.auto_filter.ref = ws_data.dimensions

        for col in ws_data.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            col_letter = get_column_letter(col[0].column)
            ws_data.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

        wb.active = ws_data
        return wb

        wb.active = ws_data
        return wb

    def create_reference_matched_report_workbook(self, p: Any) -> openpyxl.Workbook:
        """
        Creates a multi-sheet, reference-matched Excel workbook:
        Sheet 1: 'Overview' (Branded dark navy banner, 4 KPI cards with pastel fills,
                 openpyxl DoughnutChart and ColumnChart, Category-wise / Summary table with formulas,
                 Product Details preview, Low Stock alert box, Valuation Summary)
        Sheet 2+: Dedicated Detail worksheets (e.g. Product Details, Category Wise, Low Stock, Raw Data)
                 with dark navy headers, frozen header rows, auto-filters, zebra striping,
                 currency formatting (₹#,##0.00), and summary total rows with =SUM(...) formulas.
        """
        wb = openpyxl.Workbook()
        ws_overview = wb.active
        ws_overview.title = "Overview"
        ws_overview.views.sheetView[0].showGridLines = True

        fill_navy_banner = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
        fill_navy_sec = PatternFill(start_color="103B66", end_color="103B66", fill_type="solid")
        fill_teal_sec = PatternFill(start_color="0D5C3A", end_color="0D5C3A", fill_type="solid")
        fill_total_row = PatternFill(start_color="EBF3FC", end_color="EBF3FC", fill_type="solid")

        font_banner_brand = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        font_banner_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        font_banner_sub = Font(name="Calibri", size=9, bold=False, color="CBD5E1")
        font_sec_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_total = Font(name="Calibri", size=10, bold=True, color="0F2942")

        thin_side = Side(style="thin", color=self.BORDER_COLOR)
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 1. HEADER BANNER (Rows 1 to 3 across columns A to H)
        biz_name = getattr(p.branding, "business_name", None) or "Retail Store"
        ws_overview.row_dimensions[1].height = 22
        ws_overview.row_dimensions[2].height = 26
        ws_overview.row_dimensions[3].height = 6

        # Row 1: Brand (A-D) & Store/Period (E-H)
        ws_overview.merge_cells("A1:D1")
        c_brand = ws_overview.cell(row=1, column=1, value=biz_name)
        c_brand.font = font_banner_brand
        c_brand.alignment = Alignment(horizontal="left", vertical="center")

        ws_overview.merge_cells("E1:H1")
        store_period_str = f"Store: {p.store_info}  |  Period: {p.period_str}"
        c_meta = ws_overview.cell(row=1, column=5, value=store_period_str)
        c_meta.font = font_banner_sub
        c_meta.alignment = Alignment(horizontal="right", vertical="center")

        # Row 2: Title (A-H)
        ws_overview.merge_cells("A2:H2")
        c_title = ws_overview.cell(row=2, column=1, value=p.title)
        c_title.font = font_banner_title
        c_title.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx in range(1, 4):
            for c_idx in range(1, 9):
                ws_overview.cell(row=r_idx, column=c_idx).fill = fill_navy_banner

        # 2. FOUR KPI CARDS (Rows 5 to 7, Columns A to H)
        ws_overview.row_dimensions[4].height = 10
        ws_overview.row_dimensions[5].height = 16
        ws_overview.row_dimensions[6].height = 28
        ws_overview.row_dimensions[7].height = 8

        tint_config = {
            "blue": {"bg": "EBF3FC", "border": "CCE0F8", "lbl": "1E40AF"},
            "green": {"bg": "EAF7EE", "border": "C6E7CE", "lbl": "166534"},
            "orange": {"bg": "FDF3EA", "border": "FADCC3", "lbl": "C2410C"},
            "purple": {"bg": "F3EFFC", "border": "DED2F8", "lbl": "6B21A8"},
        }
        card_cols = [(1, 2), (3, 4), (5, 6), (7, 8)]
        default_tints = ["blue", "green", "orange", "purple"]

        for idx, card in enumerate(p.kpi_cards[:4]):
            c_start, c_end = card_cols[idx]
            t_key = card.get("tint") or default_tints[idx % 4]
            t_cfg = tint_config.get(t_key, tint_config["blue"])
            c_fill = PatternFill(start_color=t_cfg["bg"], end_color=t_cfg["bg"], fill_type="solid")
            c_border_side = Side(style="thin", color=t_cfg["border"])
            c_card_border = Border(left=c_border_side, right=c_border_side, top=c_border_side, bottom=c_border_side)

            ws_overview.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_end)
            ws_overview.merge_cells(start_row=6, start_column=c_start, end_row=6, end_column=c_end)

            c_lbl = ws_overview.cell(row=5, column=c_start, value=card.get("label", ""))
            c_lbl.font = Font(name="Calibri", size=8.5, bold=True, color=t_cfg["lbl"])
            c_lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            raw_val = card.get("raw_value")
            disp_val = card.get("value", "")
            c_val = ws_overview.cell(row=6, column=c_start)

            if isinstance(raw_val, (int, float, Decimal)):
                c_val.value = float(raw_val) if isinstance(raw_val, (float, Decimal)) else raw_val
                lbl_text = card.get("label", "").upper()
                if "VALUATION" in lbl_text or "SALES" in lbl_text or "REVENUE" in lbl_text or "PROFIT" in lbl_text:
                    c_val.number_format = "₹#,##0.00"
                else:
                    c_val.number_format = "#,##0"
            else:
                c_val.value = disp_val

            c_val.font = Font(name="Calibri", size=16, bold=True, color="0F2942")
            c_val.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            for r_card in range(5, 8):
                for col_card in range(c_start, c_end + 1):
                    cell = ws_overview.cell(row=r_card, column=col_card)
                    cell.fill = c_fill
                    cell.border = c_card_border

        ws_overview.row_dimensions[8].height = 10
        ws_overview.row_dimensions[9].height = 10

        # 3. OVERVIEW SUMMARY TABLES (Rows 21+)
        current_row = 21
        cat_data_start_row = 0
        cat_data_end_row = 0

        if p.summary_table and p.summary_table.get("rows"):
            s_tab = p.summary_table
            ws_overview.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(s_tab["headers"]))
            c_sec = ws_overview.cell(row=current_row, column=1, value=s_tab.get("title", "Category-wise Valuation"))
            c_sec.font = font_sec_hdr
            c_sec.fill = fill_teal_sec
            c_sec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            for c_idx, h in enumerate(s_tab["headers"], start=1):
                cell = ws_overview.cell(row=current_row, column=c_idx, value=h)
                cell.font = font_sec_hdr
                cell.fill = fill_teal_sec
                cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
                cell.border = cell_border
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            cat_data_start_row = current_row

            for r in s_tab["rows"]:
                ws_overview.row_dimensions[current_row].height = 20
                for c_idx, val in enumerate(r, start=1):
                    cell = ws_overview.cell(row=current_row, column=c_idx)
                    cell.font = self.data_font
                    cell.border = cell_border
                    num_val, fmt = self._parse_numeric(val)
                    if num_val is not None:
                        cell.value = num_val
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = fmt
                    else:
                        cell.value = val
                        cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
                current_row += 1

            cat_data_end_row = current_row - 1

            if s_tab.get("total_row"):
                ws_overview.row_dimensions[current_row].height = 22
                for c_idx in range(1, len(s_tab["headers"]) + 1):
                    cell = ws_overview.cell(row=current_row, column=c_idx)
                    cell.font = font_total
                    cell.fill = fill_total_row
                    cell.border = cell_border
                    if c_idx == 1:
                        cell.value = "Total"
                    elif c_idx == 2:
                        cell.value = ""
                    elif c_idx in (3, 4):
                        col_letter = get_column_letter(c_idx)
                        cell.value = f"=SUM({col_letter}{cat_data_start_row}:{col_letter}{cat_data_end_row})"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0"
                    elif c_idx in (5, 6):
                        col_letter = get_column_letter(c_idx)
                        cell.value = f"=SUM({col_letter}{cat_data_start_row}:{col_letter}{cat_data_end_row})"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "₹#,##0.00"
                    elif c_idx == 7:
                        cell.value = "100.00%"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                current_row += 1

            current_row += 2

        # 4. CHARTS ON OVERVIEW SHEET (Rows 10 to 19)
        if cat_data_start_row > 0 and cat_data_end_row >= cat_data_start_row:
            try:
                from openpyxl.chart import DoughnutChart, BarChart, Reference

                doughnut = DoughnutChart()
                doughnut.title = "Retail Value by Category"
                doughnut.width = 11.5
                doughnut.height = 6.8

                data_ref = Reference(ws_overview, min_col=6, min_row=cat_data_start_row - 1, max_row=cat_data_end_row)
                cats_ref = Reference(ws_overview, min_col=2, min_row=cat_data_start_row, max_row=cat_data_end_row)
                doughnut.add_data(data_ref, titles_from_data=True)
                doughnut.set_categories(cats_ref)
                ws_overview.add_chart(doughnut, "A10")

                col_chart = BarChart()
                col_chart.type = "col"
                col_chart.title = "Top Retail Values"
                col_chart.width = 11.5
                col_chart.height = 6.8
                col_chart.legend = None
                col_chart.add_data(data_ref, titles_from_data=True)
                col_chart.set_categories(cats_ref)
                ws_overview.add_chart(col_chart, "E10")
            except Exception as e:
                logger.warning(f"Failed to attach openpyxl charts on Overview sheet: {e}")

        # Detail Table preview on Overview
        if p.detail_table and p.detail_table.get("rows"):
            d_tab = p.detail_table
            ws_overview.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(d_tab["headers"]))
            c_sec = ws_overview.cell(row=current_row, column=1, value=d_tab.get("title", "Product Details"))
            c_sec.font = font_sec_hdr
            c_sec.fill = fill_navy_sec
            c_sec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            for c_idx, h in enumerate(d_tab["headers"], start=1):
                cell = ws_overview.cell(row=current_row, column=c_idx, value=h)
                cell.font = font_sec_hdr
                cell.fill = fill_navy_sec
                cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
                cell.border = cell_border
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            det_start_row = current_row
            for r in d_tab["rows"][:10]:
                ws_overview.row_dimensions[current_row].height = 20
                for c_idx, val in enumerate(r, start=1):
                    cell = ws_overview.cell(row=current_row, column=c_idx)
                    cell.font = self.data_font
                    cell.border = cell_border
                    num_val, fmt = self._parse_numeric(val)
                    if num_val is not None:
                        cell.value = num_val
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = fmt
                    else:
                        cell.value = val
                        cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
                current_row += 1

            det_end_row = current_row - 1

            if d_tab.get("total_row"):
                ws_overview.row_dimensions[current_row].height = 22
                for c_idx in range(1, len(d_tab["headers"]) + 1):
                    cell = ws_overview.cell(row=current_row, column=c_idx)
                    cell.font = font_total
                    cell.fill = fill_total_row
                    cell.border = cell_border
                    if c_idx == 1:
                        cell.value = "Total"
                    elif c_idx == 6:
                        col_letter = get_column_letter(c_idx)
                        cell.value = f"=SUM({col_letter}{det_start_row}:{col_letter}{det_end_row})"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0"
                    elif c_idx == 9:
                        col_letter = get_column_letter(c_idx)
                        cell.value = f"=SUM({col_letter}{det_start_row}:{col_letter}{det_end_row})"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "₹#,##0.00"
                current_row += 1

            current_row += 2

        # Low Stock section on Overview
        if p.exception_section:
            ex_sec = p.exception_section
            ws_overview.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
            c_sec = ws_overview.cell(row=current_row, column=1, value=ex_sec.get("title", "Low Stock Items"))
            c_sec.font = font_sec_hdr
            c_sec.fill = PatternFill(start_color="1E5086", end_color="1E5086", fill_type="solid")
            c_sec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            if not ex_sec.get("rows"):
                ws_overview.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
                c_msg = ws_overview.cell(row=current_row, column=1, value=ex_sec.get("empty_message", "No low stock items found."))
                c_msg.font = Font(name="Calibri", size=10, bold=True, color="166534")
                c_msg.fill = PatternFill(start_color="EAF7EE", end_color="EAF7EE", fill_type="solid")
                c_msg.alignment = Alignment(horizontal="center", vertical="center")
                ws_overview.row_dimensions[current_row].height = 24
                current_row += 2
            else:
                for c_idx, h in enumerate(ex_sec["headers"][:6], start=1):
                    cell = ws_overview.cell(row=current_row, column=c_idx, value=h)
                    cell.font = font_sec_hdr
                    cell.fill = PatternFill(start_color="1E5086", end_color="1E5086", fill_type="solid")
                    cell.border = cell_border
                current_row += 1
                for r in ex_sec["rows"]:
                    for c_idx, val in enumerate(r[:6], start=1):
                        cell = ws_overview.cell(row=current_row, column=c_idx, value=val)
                        cell.font = self.data_font
                        cell.border = cell_border
                    current_row += 1
                current_row += 2

        # Valuation Summary table on Overview
        if p.valuation_summary:
            ws_overview.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            c_sec = ws_overview.cell(row=current_row, column=1, value="Valuation Summary")
            c_sec.font = font_sec_hdr
            c_sec.fill = fill_navy_sec
            ws_overview.row_dimensions[current_row].height = 22
            current_row += 1

            for item in p.valuation_summary:
                c1 = ws_overview.cell(row=current_row, column=1, value=item["metric"])
                c2 = ws_overview.cell(row=current_row, column=2, value=item["amount"])
                c1.font = Font(name="Calibri", size=10, bold=True, color="1F2937")
                c2.font = self.data_font
                c1.border = cell_border
                c2.border = cell_border
                num_v, fmt_v = self._parse_numeric(item["amount"])
                if num_v is not None:
                    c2.value = num_v
                    c2.number_format = fmt_v
                    c2.alignment = Alignment(horizontal="right", vertical="center")
                current_row += 1

        # Adjust Overview Column Widths
        col_widths_overview = {1: 8, 2: 24, 3: 14, 4: 14, 5: 18, 6: 18, 7: 14, 8: 18}
        for col_idx, width in col_widths_overview.items():
            col_letter = get_column_letter(col_idx)
            ws_overview.column_dimensions[col_letter].width = width

        # -------------------------------------------------------------
        # 5. MULTI-SHEET DETAIL TABS
        # -------------------------------------------------------------
        for ws_spec in p.worksheets:
            if ws_spec.get("is_dashboard"):
                continue

            tab_title = ws_spec.get("title", "Details")[:31].replace(":", "").replace("/", "_").replace("\\", "_")
            ws_tab = wb.create_sheet(title=tab_title)
            ws_tab.freeze_panes = "A2"
            ws_tab.views.sheetView[0].showGridLines = True

            t_headers = ws_spec.get("headers", [])
            t_rows = ws_spec.get("rows", [])

            ws_tab.append(t_headers)
            ws_tab.row_dimensions[1].height = 26
            for col_idx in range(1, len(t_headers) + 1):
                cell = ws_tab.cell(row=1, column=col_idx)
                cell.font = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
                cell.fill = fill_navy_sec
                cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left", vertical="center")
                cell.border = cell_border

            start_data_r = 2
            for row_idx, r in enumerate(t_rows, start=2):
                ws_tab.row_dimensions[row_idx].height = 20
                is_even = (row_idx % 2 == 0)
                fill_row = self.zebra_fill if not is_even else None

                for col_idx, val in enumerate(r, start=1):
                    cell = ws_tab.cell(row=row_idx, column=col_idx)
                    cell.font = self.data_font
                    cell.border = cell_border
                    if fill_row:
                        cell.fill = fill_row

                    num_val, fmt = self._parse_numeric(val)
                    if num_val is not None:
                        cell.value = num_val
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = fmt
                    else:
                        cell.value = val
                        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left", vertical="center")

            end_data_r = len(t_rows) + 1

            if ws_spec.get("has_total") and len(t_rows) > 0:
                tot_r_idx = end_data_r + 1
                ws_tab.row_dimensions[tot_r_idx].height = 22
                for col_idx in range(1, len(t_headers) + 1):
                    cell = ws_tab.cell(row=tot_r_idx, column=col_idx)
                    cell.font = font_total
                    cell.fill = fill_total_row
                    cell.border = cell_border
                    if col_idx == 1:
                        cell.value = "Total"
                    elif col_idx in ws_spec.get("total_cols", []):
                        col_letter = get_column_letter(col_idx)
                        cell.value = f"=SUM({col_letter}{start_data_r}:{col_letter}{end_data_r})"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        hdr_text = t_headers[col_idx - 1].lower()
                        if "₹" in hdr_text or "value" in hdr_text or "price" in hdr_text or "amount" in hdr_text or "sales" in hdr_text or "cost" in hdr_text or "revenue" in hdr_text:
                            cell.number_format = "₹#,##0.00"
                        elif "share" in hdr_text or "%" in hdr_text:
                            cell.number_format = "0.00%"
                        else:
                            cell.number_format = "#,##0"

            if len(t_headers) > 0 and len(t_rows) > 0:
                ws_tab.auto_filter.ref = f"A1:{get_column_letter(len(t_headers))}{end_data_r}"

            for col in ws_tab.columns:
                max_len = max((len(str(c.value or "")) for c in col), default=10)
                col_letter = get_column_letter(col[0].column)
                ws_tab.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        detail_sheets = [s for s in wb.worksheets if s.title != "Overview"]
        if detail_sheets:
            wb.active = detail_sheets[0]
        else:
            wb.active = ws_overview
        return wb

    def _parse_numeric(self, val: Any) -> Tuple[Optional[float], Optional[str]]:
        if val is None or val == "":
            return None, None
        if isinstance(val, (int, float, Decimal)):
            f = float(val)
            return f, ("#,##0.00" if isinstance(val, (float, Decimal)) else "#,##0")
        if isinstance(val, str):
            s = val.strip().replace(",", "")
            if s.endswith("%"):
                try:
                    f = float(s.rstrip("%")) / 100.0
                    return f, "0.00%"
                except ValueError:
                    pass
            if s.startswith("₹"):
                try:
                    f = float(s.lstrip("₹"))
                    return f, "₹#,##0.00"
                except ValueError:
                    pass
            try:
                f = float(s)
                return f, ("#,##0.00" if "." in s else "#,##0")
            except ValueError:
                pass
        return None, None

    def create_report_streaming_response(
        self,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
        filename: str,
        metadata: dict[str, Any] | None = None,
        kpis: list[dict[str, Any]] | None = None,
        branding: Any = None,
        enable_zebra: bool = True,
        report_type: str | None = None,
        presentation: Any = None,
    ) -> StreamingResponse:
        """Generates a professional reference-matched multi-sheet report workbook and returns a StreamingResponse."""
        clean_name = filename if filename.endswith(".xlsx") else f"{filename}.xlsx"
        if presentation is not None:
            wb = self.create_reference_matched_report_workbook(presentation)
        else:
            wb = self.create_report_workbook(
                sheet_title=sheet_title,
                headers=headers,
                rows=rows,
                metadata=metadata,
                kpis=kpis,
                branding=branding,
                enable_zebra=enable_zebra,
                report_type=report_type,
            )
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={clean_name}"},
        )

