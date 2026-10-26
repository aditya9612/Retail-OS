from typing import Optional
import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import load_workbook
from sqlalchemy.orm.attributes import flag_modified

from app.services.excel_export_service import ExcelExportService

from app.core.exceptions import AppException, ConflictException, NotFoundException
from app.models.delivery import Delivery
from app.models.delivery_method import DeliveryMethod
from app.models.delivery_zone import DeliveryZone
from app.models.delivery_partner import DeliveryPartner
from app.models.order import Order, OrderTracking
from app.repositories.delivery_repo import DeliveryRepository
from app.schemas.delivery import (
    DeliveryLabelData,
    DeliveryMethodCreate,
    DeliveryMethodUpdate,
    DeliveryPartnerConnect,
    DeliveryPartnerUpdate,
    DeliveryStatsData,
    DeliveryStatsResponse,
    DeliveryTrackingData,
    DeliveryZoneCreate,
    DeliveryZoneUpdate,
    ServiceabilityData,
    ServiceabilityUploadData,
)
from app.utils.barcode_generator import generate_barcode


DELIVERY_STATUS_TRANSITIONS: dict[str, set[str]] = {
    "pending": {"assigned", "out_for_delivery", "delivered", "cancelled"},
    "assigned": {"out_for_delivery", "delivered", "cancelled"},
    "out_for_delivery": {"delivered", "cancelled"},
    "delivered": set(),
    "cancelled": set(),
}


class DeliveryService:

    def __init__(self, db):
        self.db = db
        self.repo = DeliveryRepository(db)

    def list_deliveries(
        self,
        tenant_id: int,
        page: int | None = None,
        page_size: int | None = None,
        status: str | None = None,
        delivery_person: str | None = None,
        partner_id: int | None = None,
        order_id: int | None = None,
        is_cod: bool | None = None,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
        search: str | None = None,
        store_id: int | None = None,
    ) -> tuple[list[Delivery], int]:
        return self.repo.list_deliveries(
            tenant_id=tenant_id,
            page=page,
            page_size=page_size,
            status=status,
            delivery_person=delivery_person,
            partner_id=partner_id,
            order_id=order_id,
            is_cod=is_cod,
            date_from=date_from,
            date_to=date_to,
            search=search,
            store_id=store_id,
        )

    def get_delivery(self, tenant_id: int, delivery_id: int):
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.repo.get_by_id(delivery_id, tenant_id)

        if not delivery:
            raise NotFoundException("Delivery not found")

        return delivery

    def update_status(
        self,
        tenant_id: int,
        delivery_id: int,
        status: str,
    ):
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(
            tenant_id,
            delivery_id,
        )

        target_status = status.strip().lower()
        current_status = (
            delivery.status.strip().lower() if delivery.status else "pending"
        )

        if target_status == current_status:
            raise ConflictException(f"Delivery is already {current_status}")

        allowed_next = DELIVERY_STATUS_TRANSITIONS.get(current_status, set())
        if target_status not in allowed_next:
            raise ConflictException(
                f"Cannot change delivery status from '{current_status}' to '{target_status}'"
            )

        delivery.status = target_status

        if target_status == "delivered":
            if not delivery.delivered_at:
                delivery.delivered_at = datetime.utcnow()
            if delivery.is_cod and not delivery.cod_collected:
                delivery.cod_collected = True
                delivery.cod_collected_at = datetime.utcnow()

        if delivery.order:
            delivery.order.delivery_status = target_status

        tracking = OrderTracking(
            order_id=delivery.order_id,
            status=target_status,
            remarks=f"Order {target_status.replace('_', ' ')}",
        )
        self.db.add(tracking)

        return self.repo.update(delivery)

    def create_delivery(
        self,
        tenant_id: int,
        order_id: int,
        delivery_person: str | None = None,
        tracking_number: str | None = None,
        partner_id: int | None = None,
        delivery_method_id: int | None = None,
        zone_id: int | None = None,
        is_cod: bool | None = None,
        cod_amount: Decimal | None = None,
        delivery_slot: str | None = None,
        estimated_delivery_time: datetime | None = None,
        dispatched_at: datetime | None = None,
    ) -> Delivery:
        if not isinstance(order_id, int) or order_id <= 0:
            raise AppException(
                "Order ID must be a positive integer",
                status_code=422,
            )

        order = (
            self.db.query(Order)
            .filter(Order.id == order_id, Order.tenant_id == tenant_id)
            .first()
        )
        if not order:
            raise NotFoundException("Order not found")

        if order.delivery is not None:
            raise ConflictException(f"Delivery already exists for order {order_id}")

        order_status = (order.status or "").strip().lower()
        if order_status in {"cancelled", "returned", "refunded"}:
            raise ConflictException(
                f"Cannot create delivery for order in '{order_status}' status"
            )

        partner = None
        if partner_id is not None:
            partner = self.get_partner(tenant_id, partner_id)
            if not partner.is_active:
                raise ConflictException(f"Delivery partner '{partner.name}' is inactive")
            if not delivery_person:
                delivery_person = partner.name

        if delivery_method_id is not None:
            method = self.get_method(tenant_id, delivery_method_id)
            if not method.is_active:
                raise ConflictException(f"Delivery method '{method.name}' is inactive")

        if zone_id is not None:
            zone = self.get_zone(tenant_id, zone_id)
            if not zone.is_active:
                raise ConflictException(f"Delivery zone '{zone.name}' is inactive")

        order_is_cod = False
        if is_cod is not None:
            order_is_cod = is_cod
        elif getattr(order, "payment_status", None) == "pending":
            order_is_cod = True

        final_cod_amount = Decimal("0.00")
        if order_is_cod:
            if cod_amount is not None:
                final_cod_amount = cod_amount
            elif getattr(order, "total_amount", None) is not None:
                final_cod_amount = order.total_amount

        import secrets
        delivery_otp = f"{secrets.randbelow(900000) + 100000}"

        status = "assigned" if delivery_person else "pending"
        delivery = Delivery(
            tenant_id=tenant_id,
            order_id=order_id,
            status=status,
            delivery_person=delivery_person,
            tracking_number=tracking_number,
            partner_id=partner_id,
            delivery_method_id=delivery_method_id,
            zone_id=zone_id,
            is_cod=order_is_cod,
            cod_amount=final_cod_amount,
            delivery_slot=delivery_slot,
            estimated_delivery_time=estimated_delivery_time,
            dispatched_at=dispatched_at,
            delivery_otp=delivery_otp,
        )
        self.repo.create(delivery)

        order.delivery_status = status
        tracking = OrderTracking(
            order_id=order_id,
            status=status,
            remarks=f"Delivery created with status '{status}'" + (f" via {partner.name}" if partner else ""),
        )
        self.db.add(tracking)
        self.db.commit()
        self.db.refresh(delivery)
        return delivery

    def cancel_delivery(
        self,
        tenant_id: int,
        delivery_id: int,
        reason: str | None = None,
    ) -> Delivery:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        current_status = (delivery.status or "pending").strip().lower()

        if current_status == "cancelled":
            raise ConflictException("Delivery is already cancelled")

        if current_status == "delivered":
            raise ConflictException("Cannot cancel a delivered shipment")

        allowed_next = DELIVERY_STATUS_TRANSITIONS.get(current_status, set())
        if "cancelled" not in allowed_next:
            raise ConflictException(
                f"Cannot cancel delivery in '{current_status}' status"
            )

        delivery.status = "cancelled"
        if delivery.order:
            delivery.order.delivery_status = "cancelled"

        remarks = f"Delivery cancelled: {reason}" if reason else "Delivery cancelled"
        tracking = OrderTracking(
            order_id=delivery.order_id,
            status="cancelled",
            remarks=remarks,
        )
        self.db.add(tracking)

        return self.repo.update(delivery)

    def assign_partner(
        self,
        tenant_id: int,
        delivery_id: int,
        delivery_person: str | None = None,
        tracking_number: str | None = None,
        partner_id: int | None = None,
    ) -> Delivery:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        current_status = (delivery.status or "pending").strip().lower()

        if current_status in {"delivered", "cancelled"}:
            raise ConflictException(
                f"Cannot assign delivery partner to delivery in '{current_status}' status"
            )

        if partner_id is not None:
            partner = self.get_partner(tenant_id, partner_id)
            if not partner.is_active:
                raise ConflictException(f"Delivery partner '{partner.name}' is inactive")
            delivery.partner_id = partner.id
            if not delivery_person:
                delivery_person = partner.name

        if delivery_person:
            delivery.delivery_person = delivery_person

        if tracking_number:
            delivery.tracking_number = tracking_number

        if current_status == "pending":
            delivery.status = "assigned"
            if delivery.order:
                delivery.order.delivery_status = "assigned"

        assigned_name = delivery.delivery_person or "Delivery Partner"
        remarks = f"Delivery assigned to {assigned_name}"
        if tracking_number:
            remarks += f" (Tracking: {tracking_number})"
        tracking = OrderTracking(
            order_id=delivery.order_id,
            status=delivery.status,
            remarks=remarks,
        )
        self.db.add(tracking)

        return self.repo.update(delivery)

    def record_failed_attempt(
        self,
        tenant_id: int,
        delivery_id: int,
        reason: str,
        remarks: str | None = None,
        reschedule_date: datetime | None = None,
    ) -> Delivery:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        current_status = (delivery.status or "pending").strip().lower()

        if current_status in {"delivered", "cancelled"}:
            raise ConflictException(
                f"Cannot record failed delivery attempt for delivery in '{current_status}' status"
            )

        delivery.attempt_count = (delivery.attempt_count or 0) + 1
        delivery.failure_reason = reason
        if reschedule_date:
            delivery.estimated_delivery_time = reschedule_date

        msg = f"Delivery attempt #{delivery.attempt_count} failed: {reason}"
        if remarks:
            msg += f" (Remarks: {remarks})"
        if reschedule_date:
            msg += f" [Rescheduled: {reschedule_date.strftime('%Y-%m-%d %H:%M')}]"

        tracking = OrderTracking(
            order_id=delivery.order_id,
            status=delivery.status,
            remarks=msg,
        )
        self.db.add(tracking)

        return self.repo.update(delivery)

    def get_thermal_label_html(
        self,
        tenant_id: int,
        delivery_id: int,
        auto_print: bool = True,
    ) -> str:
        delivery = self.get_delivery(tenant_id, delivery_id)
        order = delivery.order
        if not order:
            order = (
                self.db.query(Order)
                .filter(Order.id == delivery.order_id, Order.tenant_id == tenant_id)
                .first()
            )
        if not order:
            raise NotFoundException("Associated order not found")

        recipient_name = delivery.recipient_name or "Customer"
        recipient_phone = delivery.recipient_phone or "N/A"
        delivery_address = delivery.delivery_address or "No address provided"
        barcode = generate_barcode(f"{order.order_number}-{delivery.id}")
        awb_no = delivery.tracking_number or f"TRK-{order.order_number}"
        carrier_name = delivery.partner_name or delivery.delivery_person or "Standard Fleet"
        date_str = delivery.created_at.strftime("%d-%b-%Y")
        cod_badge = f'<div class="badge cod">COD: ₹{delivery.cod_amount:.2f}</div>' if delivery.is_cod else '<div class="badge prepaid">PREPAID</div>'

        body_tag = '<body onload="window.print()">' if auto_print else '<body>'
        html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Shipping Label - {order.order_number}</title>
<style>
  @page {{ size: 4in 6in; margin: 0; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
    margin: 0;
    padding: 10px;
    background: #fff;
    color: #111;
    width: 3.8in;
    box-sizing: border-box;
  }}
  .label-container {{
    border: 2px solid #000;
    padding: 8px;
    height: 5.7in;
    box-sizing: border-box;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
  }}
  .header {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    border-bottom: 2px solid #000;
    padding-bottom: 6px;
  }}
  .carrier-title {{
    font-size: 16px;
    font-weight: 800;
    text-transform: uppercase;
  }}
  .badge {{
    padding: 4px 8px;
    font-size: 12px;
    font-weight: 800;
    border: 2px solid #000;
    border-radius: 4px;
    text-align: center;
  }}
  .badge.cod {{ background: #000; color: #fff; }}
  .badge.prepaid {{ background: #fff; color: #000; }}
  .barcode-section {{
    text-align: center;
    padding: 8px 0;
    border-bottom: 1px dashed #666;
  }}
  .barcode-img {{
    max-width: 90%;
    height: 50px;
  }}
  .awb-text {{
    font-family: monospace;
    font-size: 12px;
    font-weight: bold;
    margin-top: 3px;
    letter-spacing: 1px;
  }}
  .address-section {{
    border-bottom: 1px solid #000;
    padding: 6px 0;
  }}
  .sec-title {{
    font-size: 9px;
    text-transform: uppercase;
    color: #555;
    font-weight: bold;
    margin-bottom: 2px;
  }}
  .recipient-name {{
    font-size: 14px;
    font-weight: bold;
  }}
  .address-text {{
    font-size: 11px;
    line-height: 1.3;
    margin: 3px 0;
  }}
  .phone-text {{
    font-size: 11px;
    font-weight: bold;
  }}
  .footer-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 6px;
    padding-top: 6px;
    border-top: 2px solid #000;
    font-size: 10px;
  }}
  .footer-cell {{
    display: flex;
    flex-direction: column;
  }}
</style>
</head>
{body_tag}
  <div class="label-container">
    <div class="header">
      <div class="carrier-title">{carrier_name}</div>
      {cod_badge}
    </div>

    <div class="barcode-section">
      <img class="barcode-img" src="data:image/png;base64,{barcode}" alt="AWB Barcode" />
      <div class="awb-text">{awb_no}</div>
    </div>

    <div class="address-section">
      <div class="sec-title">SHIP TO / DELIVER TO:</div>
      <div class="recipient-name">{recipient_name}</div>
      <div class="address-text">{delivery_address}</div>
      <div class="phone-text">Contact: {recipient_phone}</div>
    </div>

    <div class="footer-grid">
      <div class="footer-cell">
        <span class="sec-title">ORDER #</span>
        <strong>{order.order_number}</strong>
      </div>
      <div class="footer-cell">
        <span class="sec-title">DISPATCH DATE</span>
        <span>{date_str}</span>
      </div>
      <div class="footer-cell">
        <span class="sec-title">DELIVERY ID</span>
        <span>#{delivery.id}</span>
      </div>
      <div class="footer-cell">
        <span class="sec-title">STATUS</span>
        <span>{delivery.status.upper()}</span>
      </div>
    </div>
  </div>
</body>
</html>"""
        return html

    def get_thermal_label_image(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> bytes:
        from app.utils.barcode_generator import generate_thermal_shipping_label_image

        delivery = self.get_delivery(tenant_id, delivery_id)
        order = delivery.order
        if not order:
            order = (
                self.db.query(Order)
                .filter(Order.id == delivery.order_id, Order.tenant_id == tenant_id)
                .first()
            )
        if not order:
            raise NotFoundException("Associated order not found")

        recipient_name = delivery.recipient_name or "Retail Customer"
        recipient_phone = delivery.recipient_phone or "N/A"
        delivery_address = delivery.delivery_address or "No address provided"
        barcode_val = generate_barcode(f"{order.order_number}-{delivery.id}")
        awb_no = delivery.tracking_number or f"TRK-{order.order_number}"
        carrier_name = delivery.partner_name or delivery.delivery_person or "Standard Fleet"
        date_str = (
            delivery.dispatched_at.strftime("%d-%b-%Y")
            if delivery.dispatched_at
            else (delivery.created_at.strftime("%d-%b-%Y") if delivery.created_at else "N/A")
        )
        cod_amount = float(delivery.cod_amount or 0.0)

        return generate_thermal_shipping_label_image(
            carrier_name=carrier_name,
            is_cod=bool(delivery.is_cod),
            cod_amount=cod_amount,
            awb_number=awb_no,
            barcode_data=barcode_val,
            recipient_name=recipient_name,
            address=delivery_address,
            phone=recipient_phone,
            order_number=order.order_number,
            dispatch_date=date_str,
            delivery_id=delivery.id,
            status=delivery.status or "PENDING",
        )

    def get_thermal_label_pdf(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> bytes:
        img_bytes = self.get_thermal_label_image(tenant_id, delivery_id)
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import inch
        from reportlab.lib.utils import ImageReader
        import io

        pdf_buf = io.BytesIO()
        c = canvas.Canvas(pdf_buf, pagesize=(4 * inch, 6 * inch))
        c.drawImage(ImageReader(io.BytesIO(img_bytes)), 0, 0, width=4 * inch, height=6 * inch)
        c.showPage()
        c.save()
        return pdf_buf.getvalue()

    def get_thermal_label_text(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> str:
        delivery = self.get_delivery(tenant_id, delivery_id)
        order = delivery.order
        if not order and delivery.order_id:
            from app.models.order import Order

            order = (
                self.db.query(Order)
                .filter(Order.id == delivery.order_id, Order.tenant_id == tenant_id)
                .first()
            )
        order_number = (
            order.order_number
            if order and order.order_number
            else f"ORD-{delivery.order_id}"
        )
        awb_no = delivery.tracking_number or f"TRK-{order_number}"
        carrier_name = delivery.partner_name or delivery.delivery_person or "Standard Fleet"
        recipient_name = delivery.recipient_name or "Retail Customer"
        recipient_phone = delivery.recipient_phone or "N/A"
        delivery_address = delivery.delivery_address or "No address provided"
        cod_str = f"COD: Rs. {delivery.cod_amount:.2f}" if delivery.is_cod else "PREPAID"
        dt_str = (
            delivery.dispatched_at.strftime("%Y-%m-%d %H:%M")
            if delivery.dispatched_at
            else (
                delivery.created_at.strftime("%Y-%m-%d %H:%M")
                if delivery.created_at
                else ""
            )
        )

        sep = "=" * 32
        dashes = "-" * 32
        lines = [
            sep,
            "SHIPPING LABEL (4x6 / 80mm)".center(32),
            sep,
            f"CARRIER : {carrier_name}",
            f"AWB NO  : {awb_no}",
            f"ORDER NO: {order_number}",
            f"TYPE    : {cod_str}",
            dashes,
            "SHIP TO:",
            f"  {recipient_name}",
            f"  Tel: {recipient_phone}",
            f"  {delivery_address[:60]}",
            dashes,
            f"DATE    : {dt_str}",
            sep,
            f"BARCODE : *{awb_no}*",
            sep,
        ]
        return "\n".join(lines)

    def collect_cod(
        self,
        tenant_id: int,
        delivery_id: int,
        amount_collected: Decimal | None = None,
    ) -> Delivery:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        if not delivery.is_cod:
            raise ConflictException("Delivery is not marked as Cash On Delivery (COD)")

        if delivery.cod_collected:
            raise ConflictException("COD cash has already been collected for this delivery")

        delivery.cod_collected = True
        delivery.cod_collected_at = datetime.utcnow()
        if amount_collected is not None:
            delivery.cod_amount = amount_collected

        tracking = OrderTracking(
            order_id=delivery.order_id,
            status=delivery.status,
            remarks=f"COD payment collected: ₹{delivery.cod_amount:.2f}",
        )
        self.db.add(tracking)
        return self.repo.update(delivery)

    def get_cod_reconciliation(
        self,
        tenant_id: int,
        delivery_person: str | None = None,
        partner_id: int | None = None,
        status: str | None = None,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
        store_id: Optional[int] = None,
    ) -> dict:
        return self.repo.get_cod_reconciliation(
            tenant_id=tenant_id,
            delivery_person=delivery_person,
            partner_id=partner_id,
            status=status,
            date_from=date_from,
            date_to=date_to,
            store_id=store_id,
        )

    def update_address(
        self,
        tenant_id: int,
        delivery_id: int,
        delivery_address: str,
    ) -> Delivery:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        current_status = (delivery.status or "pending").strip().lower()

        if current_status in {"delivered", "cancelled"}:
            raise ConflictException(
                f"Cannot update delivery address for delivery in '{current_status}' status"
            )

        order = delivery.order
        if not order:
            order = (
                self.db.query(Order)
                .filter(Order.id == delivery.order_id, Order.tenant_id == tenant_id)
                .first()
            )
        if not order:
            raise NotFoundException("Associated order not found")

        order.delivery_address = delivery_address

        tracking = OrderTracking(
            order_id=delivery.order_id,
            status=delivery.status,
            remarks=f"Delivery address updated: {delivery_address}",
        )
        self.db.add(tracking)

        return self.repo.update(delivery)

    def get_label(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> DeliveryLabelData:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        order = delivery.order
        if not order:
            order = (
                self.db.query(Order)
                .filter(Order.id == delivery.order_id, Order.tenant_id == tenant_id)
                .first()
            )
        if not order:
            raise NotFoundException("Associated order not found")

        if order.customer:
            recipient_name = order.customer.name
            recipient_phone = order.customer.phone
        else:
            recipient_name = "Retail Customer"
            recipient_phone = None

        raw_barcode = generate_barcode(f"{order.order_number}-{delivery.id}")

        return DeliveryLabelData(
            delivery_id=delivery.id,
            order_id=order.id,
            order_number=order.order_number,
            tracking_number=delivery.tracking_number,
            delivery_person=delivery.delivery_person,
            status=delivery.status,
            recipient_name=recipient_name,
            recipient_phone=recipient_phone,
            delivery_address=order.delivery_address,
            barcode=raw_barcode,
            created_at=delivery.created_at,
        )

    def get_tracking(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> DeliveryTrackingData:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        latest_tracking = self.repo.get_latest_tracking(delivery.order_id)

        if latest_tracking:
            return DeliveryTrackingData(
                delivery_id=delivery.id,
                order_id=delivery.order_id,
                status=delivery.status,
                tracking_number=delivery.tracking_number,
                delivery_person=delivery.delivery_person,
                tracking_id=latest_tracking.id,
                remarks=latest_tracking.remarks,
                updated_at=latest_tracking.created_at,
            )

        return DeliveryTrackingData(
            delivery_id=delivery.id,
            order_id=delivery.order_id,
            status=delivery.status,
            tracking_number=delivery.tracking_number,
            delivery_person=delivery.delivery_person,
            tracking_id=None,
            remarks=None,
            updated_at=delivery.updated_at or delivery.created_at,
        )

    def get_history(
        self,
        tenant_id: int,
        delivery_id: int,
    ) -> list[OrderTracking]:
        if not isinstance(delivery_id, int) or delivery_id <= 0:
            raise AppException(
                "Delivery ID must be a positive integer",
                status_code=422,
            )

        delivery = self.get_delivery(tenant_id, delivery_id)
        raw_history = self.repo.get_tracking_history(delivery.order_id)
        valid_delivery_statuses = set(DELIVERY_STATUS_TRANSITIONS.keys())
        delivery_history = [
            h for h in raw_history
            if (h.status or "").strip().lower() in valid_delivery_statuses
        ]

        # If order tracking only contains order lifecycle events (e.g. "confirmed")
        # and no delivery status was logged, return matching delivery.status in-memory without database mutation
        if not delivery_history and raw_history:
            current_status = (delivery.status or "pending").strip().lower()
            return [
                OrderTracking(
                    id=raw_history[0].id,
                    order_id=delivery.order_id,
                    status=current_status,
                    remarks=f"Delivery created with status '{current_status}'",
                    created_at=delivery.created_at or datetime.utcnow(),
                )
            ]

        return delivery_history

    def get_stats(self, tenant_id: int, store_id: Optional[int] = None) -> DeliveryStatsResponse:
        stats_dict = self.repo.get_stats(tenant_id, store_id=store_id)
        data = DeliveryStatsData(**stats_dict)
        return DeliveryStatsResponse(
            success=True,
            message="Delivery statistics retrieved successfully",
            data=data,
        )

    def export_deliveries(
        self,
        tenant_id: int,
        status: str | None = None,
        format: str = "excel",
        store_id: Optional[int] = None,
    ) -> tuple[BytesIO, str, str]:
        deliveries = self.repo.get_for_export(tenant_id, status=status, store_id=store_id)

        if format == "excel":
            headers = [
                "ID",
                "Order Number",
                "Status",
                "Delivery Person",
                "Tracking Number",
                "Delivery Address",
                "Delivered At",
                "Created At",
            ]
            rows = []
            for d in deliveries:
                rows.append([
                    d.id,
                    d.order.order_number if d.order else "",
                    d.status,
                    d.delivery_person or "",
                    d.tracking_number or "",
                    (d.order.delivery_address or "") if d.order else "",
                    d.delivered_at.strftime("%Y-%m-%d %H:%M:%S") if d.delivered_at else "",
                    d.created_at.strftime("%Y-%m-%d %H:%M:%S") if d.created_at else "",
                ])
            output = ExcelExportService().export_to_stream(
                sheet_title="Deliveries",
                headers=headers,
                rows=rows,
            )
            return (
                output,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "deliveries.xlsx",
            )

        elif format == "csv":
            str_output = StringIO()
            writer = csv.writer(str_output)
            writer.writerow([
                "ID",
                "Order Number",
                "Status",
                "Delivery Person",
                "Tracking Number",
                "Delivery Address",
                "Delivered At",
                "Created At",
            ])
            for d in deliveries:
                writer.writerow([
                    d.id,
                    d.order.order_number if d.order else "",
                    d.status,
                    d.delivery_person or "",
                    d.tracking_number or "",
                    (d.order.delivery_address or "") if d.order else "",
                    d.delivered_at.strftime("%Y-%m-%d %H:%M:%S") if d.delivered_at else "",
                    d.created_at.strftime("%Y-%m-%d %H:%M:%S") if d.created_at else "",
                ])
            bytes_output = BytesIO(str_output.getvalue().encode("utf-8"))
            bytes_output.seek(0)
            return (
                bytes_output,
                "text/csv",
                "deliveries.csv",
            )

        elif format == "pdf":
            from app.services.document_settings_service import DocumentSettingsService
            from app.services.document_renderer_service import DocumentRenderer

            branding = None
            try:
                branding = DocumentSettingsService(self.db).resolve_branding(
                    tenant_id=tenant_id,
                )
            except Exception:
                branding = None

            headers = [
                "ID",
                "Order Number",
                "Status",
                "Delivery Person",
                "Tracking Number",
                "Delivery Address",
                "Delivered At",
                "Created At",
            ]
            rows = []
            for d in deliveries:
                rows.append([
                    d.id,
                    d.order.order_number if d.order else "",
                    d.status,
                    d.delivery_person or "",
                    d.tracking_number or "",
                    (d.order.delivery_address or "") if d.order else "",
                    d.delivered_at.strftime("%Y-%m-%d %H:%M:%S") if d.delivered_at else "",
                    d.created_at.strftime("%Y-%m-%d %H:%M:%S") if d.created_at else "",
                ])

            metadata = {
                "Tenant ID": str(tenant_id),
                "Status Filter": (status.capitalize() if status else "All"),
                "Total Records": str(len(deliveries)),
                "Exported At": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            }

            notes = "No delivery records found." if not rows else None

            pdf_bytes = DocumentRenderer().render_report_pdf(
                report_title="Deliveries Export",
                metadata=metadata,
                headers=headers,
                rows=rows,
                branding=branding,
                notes=notes,
                report_type="deliveries_export",
            )
            bytes_output = BytesIO(pdf_bytes)
            bytes_output.seek(0)
            return (
                bytes_output,
                "application/pdf",
                "deliveries.pdf",
            )

        raise AppException(f"Unsupported format: '{format}'", status_code=422)

    def list_methods(
        self,
        tenant_id: int,
        is_active: bool | None = None,
    ) -> list[DeliveryMethod]:
        return self.repo.list_methods(tenant_id, is_active=is_active)

    def get_method(
        self,
        tenant_id: int,
        method_id: int,
    ) -> DeliveryMethod:
        if not isinstance(method_id, int) or method_id <= 0:
            raise AppException(
                "Delivery method ID must be a positive integer",
                status_code=422,
            )

        method = self.repo.get_method_by_id(method_id, tenant_id)
        if not method:
            raise NotFoundException("Delivery method not found")

        return method

    def create_method(
        self,
        tenant_id: int,
        data: DeliveryMethodCreate,
    ) -> DeliveryMethod:
        existing = self.repo.get_method_by_code(data.code, tenant_id)
        if existing:
            raise ConflictException(
                f"Delivery method with code '{data.code}' already exists"
            )

        method = DeliveryMethod(
            tenant_id=tenant_id,
            name=data.name,
            code=data.code,
            description=data.description,
            cost=data.cost,
            estimated_days=data.estimated_days,
            is_active=data.is_active,
        )
        return self.repo.create_method(method)

    def update_method(
        self,
        tenant_id: int,
        method_id: int,
        data: DeliveryMethodUpdate,
    ) -> DeliveryMethod:
        method = self.get_method(tenant_id, method_id)

        if data.code is not None and data.code != method.code:
            existing = self.repo.get_method_by_code(data.code, tenant_id)
            if existing and existing.id != method.id:
                raise ConflictException(
                    f"Delivery method with code '{data.code}' already exists"
                )
            method.code = data.code

        if data.name is not None:
            method.name = data.name

        if data.description is not None:
            method.description = data.description

        if data.cost is not None:
            method.cost = data.cost

        if data.estimated_days is not None:
            method.estimated_days = data.estimated_days

        if data.is_active is not None:
            method.is_active = data.is_active

        return self.repo.update_method(method)

    def toggle_method(
        self,
        tenant_id: int,
        method_id: int,
    ) -> DeliveryMethod:
        method = self.get_method(tenant_id, method_id)
        method.is_active = not method.is_active
        return self.repo.update_method(method)

    def delete_method(
        self,
        tenant_id: int,
        method_id: int,
    ) -> None:
        method = self.get_method(tenant_id, method_id)
        self.repo.delete_method(method)

    def list_zones(
        self,
        tenant_id: int,
        is_active: bool | None = None,
    ) -> list[DeliveryZone]:
        return self.repo.list_zones(tenant_id, is_active=is_active)

    def get_zone(
        self,
        tenant_id: int,
        zone_id: int,
    ) -> DeliveryZone:
        if not isinstance(zone_id, int) or zone_id <= 0:
            raise AppException(
                "Delivery zone ID must be a positive integer",
                status_code=422,
            )

        zone = self.repo.get_zone_by_id(zone_id, tenant_id)
        if not zone:
            raise NotFoundException("Delivery zone not found")

        return zone

    def create_zone(
        self,
        tenant_id: int,
        data: DeliveryZoneCreate,
    ) -> DeliveryZone:
        existing = self.repo.get_zone_by_code(data.code, tenant_id)
        if existing:
            raise ConflictException(
                f"Delivery zone with code '{data.code}' already exists"
            )

        zone = DeliveryZone(
            tenant_id=tenant_id,
            name=data.name,
            code=data.code,
            description=data.description,
            city=data.city,
            state=data.state,
            pincodes=data.pincodes,
            is_active=data.is_active,
        )
        return self.repo.create_zone(zone)

    def update_zone(
        self,
        tenant_id: int,
        zone_id: int,
        data: DeliveryZoneUpdate,
    ) -> DeliveryZone:
        zone = self.get_zone(tenant_id, zone_id)

        if data.code is not None and data.code != zone.code:
            existing = self.repo.get_zone_by_code(data.code, tenant_id)
            if existing and existing.id != zone.id:
                raise ConflictException(
                    f"Delivery zone with code '{data.code}' already exists"
                )
            zone.code = data.code
        effective_state = data.state if data.state is not None else zone.state
        effective_pincodes = data.pincodes if data.pincodes is not None else zone.pincodes
        effective_city = data.city if data.city is not None else zone.city

        from app.utils.validators import validate_zone_pincodes_match_state, MAJOR_INDIAN_CITIES_STATE_MAP
        if (effective_state or effective_city) and effective_pincodes:
            try:
                validate_zone_pincodes_match_state(effective_pincodes, state=effective_state, city=effective_city)
            except ValueError as e:
                raise AppException(str(e), status_code=422)

        if effective_city and effective_state:
            c_low = effective_city.strip().lower()
            s_low = effective_state.strip().lower()
            if c_low in MAJOR_INDIAN_CITIES_STATE_MAP:
                exp_st = MAJOR_INDIAN_CITIES_STATE_MAP[c_low]
                if exp_st.lower() != s_low:
                    raise AppException(
                        f"City '{effective_city}' belongs to state '{exp_st}', not '{effective_state}'",
                        status_code=422,
                    )

        if data.name is not None:
            zone.name = data.name

        if data.description is not None:
            zone.description = data.description

        if data.city is not None:
            zone.city = data.city

        if data.state is not None:
            zone.state = data.state

        if data.pincodes is not None:
            zone.pincodes = data.pincodes

        if data.is_active is not None:
            zone.is_active = data.is_active

        return self.repo.update_zone(zone)

    def delete_zone(
        self,
        tenant_id: int,
        zone_id: int,
    ) -> None:
        zone = self.get_zone(tenant_id, zone_id)
        self.repo.delete_zone(zone)

    def list_partners(
        self,
        tenant_id: int,
        is_active: bool | None = None,
    ) -> list[DeliveryPartner]:
        return self.repo.list_partners(tenant_id, is_active=is_active)

    def get_partner(
        self,
        tenant_id: int,
        partner_id: int,
    ) -> DeliveryPartner:
        if not isinstance(partner_id, int) or partner_id <= 0:
            raise AppException(
                "Delivery partner ID must be a positive integer",
                status_code=422,
            )

        partner = self.repo.get_partner_by_id(partner_id, tenant_id)
        if not partner:
            raise NotFoundException("Delivery partner not found")

        return partner

    def create_partner(
        self,
        tenant_id: int,
        data: DeliveryPartnerConnect,
    ) -> DeliveryPartner:
        existing = self.repo.get_partner_by_code(data.code, tenant_id)
        if existing:
            raise ConflictException(
                f"Delivery partner with code '{data.code}' already exists"
            )

        partner = DeliveryPartner(
            tenant_id=tenant_id,
            name=data.name,
            code=data.code,
            description=data.description,
            contact_email=data.contact_email,
            contact_phone=data.contact_phone,
            api_key=data.api_key,
            api_secret=data.api_secret,
            tracking_url_template=data.tracking_url_template,
            is_active=data.is_active,
        )
        return self.repo.create_partner(partner)

    def update_partner(
        self,
        tenant_id: int,
        partner_id: int,
        data: DeliveryPartnerUpdate,
    ) -> DeliveryPartner:
        partner = self.get_partner(tenant_id, partner_id)

        if data.code is not None and data.code != partner.code:
            existing = self.repo.get_partner_by_code(data.code, tenant_id)
            if existing and existing.id != partner.id:
                raise ConflictException(
                    f"Delivery partner with code '{data.code}' already exists"
                )
            partner.code = data.code

        if data.name is not None:
            partner.name = data.name

        if data.description is not None:
            partner.description = data.description

        if data.contact_email is not None:
            partner.contact_email = data.contact_email

        if data.contact_phone is not None:
            partner.contact_phone = data.contact_phone

        if data.api_key is not None:
            partner.api_key = data.api_key

        if data.api_secret is not None:
            partner.api_secret = data.api_secret

        if data.tracking_url_template is not None:
            partner.tracking_url_template = data.tracking_url_template

        if data.is_active is not None:
            partner.is_active = data.is_active

        return self.repo.update_partner(partner)

    def delete_partner(
        self,
        tenant_id: int,
        partner_id: int,
    ) -> None:
        partner = self.get_partner(tenant_id, partner_id)
        self.repo.delete_partner(partner)

    def check_serviceability(
        self,
        tenant_id: int,
        pincode: str,
    ) -> ServiceabilityData:
        if not isinstance(pincode, str):
            raise AppException("Pincode must be a string", status_code=422)
        p = pincode.strip()
        if not p.isdigit() or len(p) != 6 or not (100000 <= int(p) <= 999999):
            raise AppException(
                f"Invalid pincode: '{pincode}'. Pincode must contain exactly 6 numeric digits (100000-999999)",
                status_code=422,
            )

        active_zone = self.repo.find_active_zone_by_pincode(tenant_id, p)
        if active_zone:
            return ServiceabilityData(
                pincode=p,
                is_serviceable=True,
                zone_id=active_zone.id,
                zone_name=active_zone.name,
                zone_code=active_zone.code,
                city=active_zone.city,
                state=active_zone.state,
            )

        return ServiceabilityData(
            pincode=p,
            is_serviceable=False,
            zone_id=None,
            zone_name=None,
            zone_code=None,
            city=None,
            state=None,
        )

    def upload_serviceability(
        self,
        tenant_id: int,
        filename: str,
        contents: bytes,
    ) -> ServiceabilityUploadData:
        filename_clean = (filename or "").strip().lower()
        if not (filename_clean.endswith(".csv") or filename_clean.endswith(".xlsx")):
            raise AppException(
                "Unsupported file format. Only CSV (.csv) and Excel (.xlsx) files are supported",
                status_code=422,
            )

        if not contents or len(contents) == 0:
            raise AppException("Uploaded file is empty", status_code=422)

        raw_rows: list[dict[str, Any]] = []

        if filename_clean.endswith(".csv"):
            try:
                text = contents.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = contents.decode("latin-1")
                except Exception:
                    raise AppException("Malformed or unreadable CSV file", status_code=422)

            reader = csv.DictReader(StringIO(text))
            if not reader.fieldnames:
                raise AppException("Uploaded CSV file has no header row", status_code=422)

            field_map = {name.strip().lower(): name for name in reader.fieldnames if name}
            if "pincode" not in field_map or "zone_code" not in field_map:
                raise AppException(
                    "CSV file is missing required columns. Both 'zone_code' and 'pincode' columns are required",
                    status_code=422,
                )

            for row in reader:
                if not any(v.strip() for v in row.values() if isinstance(v, str)):
                    continue
                raw_rows.append({
                    "zone_code": row.get(field_map["zone_code"]),
                    "pincode": row.get(field_map["pincode"]),
                })
        else:
            try:
                wb = load_workbook(filename=BytesIO(contents), data_only=True)
                ws = wb.active
            except Exception:
                raise AppException("Malformed or corrupted Excel file", status_code=422)

            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                raise AppException("Uploaded Excel file is empty", status_code=422)

            if not header_row:
                raise AppException("Uploaded Excel file has no header row", status_code=422)

            col_map: dict[str, int] = {}
            for idx, col_name in enumerate(header_row):
                if col_name is not None:
                    col_map[str(col_name).strip().lower()] = idx

            if "pincode" not in col_map or "zone_code" not in col_map:
                raise AppException(
                    "Excel file is missing required columns. Both 'zone_code' and 'pincode' columns are required",
                    status_code=422,
                )

            for row in rows_iter:
                if not any(row):
                    continue
                raw_rows.append({
                    "zone_code": row[col_map["zone_code"]],
                    "pincode": row[col_map["pincode"]],
                })

        if not raw_rows:
            raise AppException("Uploaded file contains no data rows", status_code=422)

        all_zones = self.repo.list_zones(tenant_id)
        zones_by_code: dict[str, DeliveryZone] = {
            z.code.strip().upper(): z for z in all_zones
        }

        db_pincode_to_zone: dict[str, str] = {}
        for z in all_zones:
            z_code = z.code.strip().upper()
            for p in (z.pincodes or []):
                db_pincode_to_zone[p] = z_code

        file_pincode_to_zone: dict[str, str] = {}
        parsed_items: list[tuple[str, str]] = []

        for row_idx, r in enumerate(raw_rows, start=2):
            raw_zone_code = r.get("zone_code")
            if raw_zone_code is None:
                raise AppException(f"Missing zone_code at row {row_idx}", status_code=422)
            zone_code = str(raw_zone_code).strip().upper()
            if not zone_code:
                raise AppException(f"Empty zone_code at row {row_idx}", status_code=422)

            if zone_code not in zones_by_code:
                raise AppException(
                    f"Delivery zone with code '{zone_code}' does not exist for this tenant (row {row_idx})",
                    status_code=422,
                )

            raw_p = r.get("pincode")
            if raw_p is None:
                raise AppException(f"Missing pincode at row {row_idx}", status_code=422)

            if isinstance(raw_p, float) or isinstance(raw_p, bool):
                raise AppException(
                    f"Invalid pincode '{raw_p}' at row {row_idx}. Pincode must be exactly 6 numeric digits (floats/booleans not allowed)",
                    status_code=422,
                )

            p_str = str(raw_p).strip()
            if not p_str.isdigit() or len(p_str) != 6 or not (100000 <= int(p_str) <= 999999):
                raise AppException(
                    f"Invalid pincode '{raw_p}' at row {row_idx}. Pincode must contain exactly 6 numeric digits (100000-999999)",
                    status_code=422,
                )

            if p_str in file_pincode_to_zone and file_pincode_to_zone[p_str] != zone_code:
                raise ConflictException(
                    f"Pincode '{p_str}' is assigned to multiple zones ('{file_pincode_to_zone[p_str]}' and '{zone_code}') in upload file (row {row_idx})"
                )
            file_pincode_to_zone[p_str] = zone_code

            if p_str in db_pincode_to_zone and db_pincode_to_zone[p_str] != zone_code:
                raise ConflictException(
                    f"Pincode '{p_str}' is already assigned to delivery zone '{db_pincode_to_zone[p_str]}'"
                )

            parsed_items.append((zone_code, p_str))

        zone_pincodes_map: dict[str, list[str]] = {
            z_code: list(z.pincodes or []) for z_code, z in zones_by_code.items()
        }

        total_rows_processed = len(parsed_items)
        pincodes_added = 0
        pincodes_existing = 0
        zones_affected: set[str] = set()

        for zone_code, p_str in parsed_items:
            current_list = zone_pincodes_map[zone_code]
            if p_str in current_list:
                pincodes_existing += 1
            else:
                current_list.append(p_str)
                pincodes_added += 1
                zones_affected.add(zone_code)

        try:
            for zone_code in zones_affected:
                zone = zones_by_code[zone_code]
                zone.pincodes = zone_pincodes_map[zone_code]
                flag_modified(zone, "pincodes")
            self.db.commit()
            for zone_code in zones_affected:
                self.db.refresh(zones_by_code[zone_code])
        except Exception as exc:
            self.db.rollback()
            raise exc

        return ServiceabilityUploadData(
            total_rows_processed=total_rows_processed,
            pincodes_added=pincodes_added,
            pincodes_existing=pincodes_existing,
            zones_updated=sorted(list(zones_affected)),
        )

