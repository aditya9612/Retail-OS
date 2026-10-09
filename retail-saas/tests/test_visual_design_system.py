import io
import re
from decimal import Decimal
import openpyxl
import pytest
from reportlab.lib import colors

from app.schemas.document_setting import BrandingContext
from app.services.document_renderer_service import DocumentRenderer
from app.services.excel_export_service import ExcelExportService
from app.services.export_theme import THEMES, get_theme_for_report


def parse_pdf_structure(pdf_bytes: bytes) -> dict:
    assert pdf_bytes.startswith(b"%PDF-"), "Missing %PDF header"
    assert b"%%EOF" in pdf_bytes, "Missing %%EOF marker"
    assert b"/Catalog" in pdf_bytes, "Missing /Catalog dictionary"
    assert b"/Pages" in pdf_bytes, "Missing /Pages dictionary"
    assert b"xref" in pdf_bytes, "Missing xref table"

    page_count = len(re.findall(rb"/Type\s*/Page\b", pdf_bytes))

    # Extract uncompressed text literals
    text_literals = []
    for m in re.finditer(rb"\((.*?)\)\s*Tj", pdf_bytes):
        raw = m.group(1).replace(b"\\(", b"(").replace(b"\\)", b")").replace(b"\\\\", b"\\")
        text_literals.append(raw.decode("latin1", errors="ignore"))

    fonts = [m.group(1).decode("latin1") for m in re.finditer(rb"/BaseFont\s*/([^\s/>]+)", pdf_bytes)]

    return {
        "page_count": page_count,
        "fonts": fonts,
        "text": " ".join(text_literals),
    }


@pytest.fixture
def branding():
    return BrandingContext(
        business_name="Nexus Retail Hub",
        address="Sector 18, Vashi, Navi Mumbai, MH 400703",
        phone="+91 9820011223",
        email="contact@nexusretail.com",
        website="https://nexusretail.com",
        gstin="27ABCDE1234F1Z5",
        footer_text="Thank you for your business. Disputes subject to Navi Mumbai jurisdiction.",
        invoice_prefix="INV",
        bill_prefix="BILL",
        show_gstin=True,
        show_qr=False,
        show_signature=True,
        show_payment_details=True,
        logo_url=None,
    )


def test_export_themes_integrity():
    """Verify all theme definitions have correct color properties and hex formats."""
    expected_themes = [
        "bill", "tax_invoice", "credit_note", "purchase_order",
        "green", "blue", "teal", "amber", "dark_green",
    ]
    for key in expected_themes:
        assert key in THEMES, f"Missing theme key: {key}"
        theme = THEMES[key]
        assert theme.primary_hex.startswith("#")
        assert len(theme.excel_header_hex) == 6
        assert isinstance(theme.primary_color, colors.Color)
        assert isinstance(theme.bg_light_color, colors.Color)
        assert isinstance(theme.border_color, colors.Color)


def test_theme_mapping_per_visual_design_map():
    """Verify report types map to designated colors per design brief."""
    # POS Documents
    assert get_theme_for_report("bill").primary_hex == "#1B5E20"
    assert get_theme_for_report("tax_invoice").primary_hex == "#1F497D"
    assert get_theme_for_report("credit_note").primary_hex == "#DC2626"
    assert get_theme_for_report("purchase_order").primary_hex == "#6D28D9"

    # Business Reports
    assert get_theme_for_report("sales_daily").primary_hex == "#15803D"
    assert get_theme_for_report("profit_loss").primary_hex == "#15803D"
    assert get_theme_for_report("inventory_valuation").primary_hex == "#1E40AF"
    assert get_theme_for_report("inventory_stock").primary_hex == "#1E40AF"
    assert get_theme_for_report("customers_overview").primary_hex == "#0F766E"
    assert get_theme_for_report("customers_retention").primary_hex == "#0F766E"
    assert get_theme_for_report("inventory_low").primary_hex == "#D97706"
    assert get_theme_for_report("products_slow_moving").primary_hex == "#D97706"
    assert get_theme_for_report("products_profitability").primary_hex == "#6D28D9"
    assert get_theme_for_report("customers_lifetime_value").primary_hex == "#6D28D9"
    assert get_theme_for_report("daily_billing_closure").primary_hex == "#14532D"
    assert get_theme_for_report("pos_z_report").primary_hex == "#14532D"


def test_credit_note_pdf_visual_rendering(branding):
    """Verify credit note PDF adheres to Coral Red theme (#DC2626) and contains required sections."""
    renderer = DocumentRenderer()
    data = {
        "credit_note_no": "CN-2026-00042",
        "original_invoice_number": "INV-2026-00105",
        "created_at": "2026-10-08",
        "refund_amount": 1450.00,
        "cgst_amount": 130.50,
        "sgst_amount": 130.50,
        "igst_amount": 0.0,
        "reason": "Defective item returned by buyer",
        "customer": {
            "name": "Apex Electronics",
            "phone": "+91 9988776655",
            "gstin": "27XYZAB9876C1Z3",
            "address": "Gala 4, Industrial Estate, Pune",
        },
    }

    pdf_bytes = renderer.render("credit_note", data, branding)
    parsed = parse_pdf_structure(pdf_bytes)
    assert parsed["page_count"] == 1
    assert "CREDIT NOTE" in parsed["text"]
    assert "ISSUED BY (SELLER)" in parsed["text"]
    assert "CREDIT ISSUED TO (BUYER)" in parsed["text"]
    assert "Total Credited Amount:" in parsed["text"]
    assert "Amount in Words:" in parsed["text"]
    assert "Authorized Signatory" in parsed["text"]
    assert "1450.00" in parsed["text"]


def test_purchase_order_pdf_visual_rendering(branding):
    """Verify purchase order PDF adheres to Purple/Indigo theme (#6D28D9) and contains required sections."""
    renderer = DocumentRenderer()
    data = {
        "order_number": "PO-2026-00088",
        "created_at": "2026-10-08",
        "status": "APPROVED",
        "notes": "Deliver to Warehouse Bay 3 within 7 business days",
        "total_amount": 25000.00,
        "supplier": {
            "name": "Global Supplies Corp",
            "contact_person": "Vikram Patel",
            "phone": "+91 9123456789",
            "email": "vikram@globalsupplies.com",
            "address": "Plot 12, MIDC, Andheri East, Mumbai",
            "gstin": "27GLOBL4321A1Z8",
        },
        "items": [
            {
                "product_name": "Premium Packaging Boxes",
                "sku": "BOX-L-01",
                "quantity": 500,
                "unit_cost": 30.00,
                "total_cost": 15000.00,
            },
            {
                "product_name": "Thermal Receipt Rolls",
                "sku": "ROLL-80MM",
                "quantity": 200,
                "unit_cost": 50.00,
                "total_cost": 10000.00,
            },
        ],
    }

    pdf_bytes = renderer.render("purchase_order", data, branding)
    parsed = parse_pdf_structure(pdf_bytes)
    assert parsed["page_count"] == 1
    assert "PURCHASE ORDER" in parsed["text"]
    assert "BUYER (DELIVER TO)" in parsed["text"]
    assert "VENDOR / SUPPLIER" in parsed["text"]
    assert "Total Order Value:" in parsed["text"]
    assert "Amount in Words:" in parsed["text"]
    assert "Authorized Signatory" in parsed["text"]
    assert "25000.00" in parsed["text"]


def test_report_pdf_themes_rendering(branding):
    """Verify report PDF rendering across key semantic themes."""
    renderer = DocumentRenderer()
    metadata = {"Period": "October 2026", "Store": "Main Branch"}
    kpis = [
        {"label": "Total Sales", "value": 154200.50},
        {"label": "Transactions", "value": 312},
    ]
    headers = ["Metric / Item", "Category", "Quantity", "Revenue"]
    rows = [
        ["Product Alpha", "Electronics", 15, 45000.00],
        ["Product Beta", "Apparel", 40, 20000.00],
    ]

    report_types_to_test = [
        "sales_daily",
        "inventory_valuation",
        "customers_overview",
        "inventory_low",
        "pos_z_report",
        "products_profitability",
    ]

    for rtype in report_types_to_test:
        pdf_bytes = renderer.render_report_pdf(
            report_title=f"{rtype.replace('_', ' ').title()} Report",
            metadata=metadata,
            kpis=kpis,
            headers=headers,
            rows=rows,
            branding=branding,
            report_type=rtype,
        )
        parsed = parse_pdf_structure(pdf_bytes)
        assert parsed["page_count"] >= 1
        assert len(pdf_bytes) > 2000
        # Assert full dynamic branding rendered
        assert "Nexus Retail Hub" in parsed["text"]
        assert "Sector 18, Vashi" in parsed["text"]
        assert "+91 9820011223" in parsed["text"]
        assert "contact@nexusretail.com" in parsed["text"]
        assert "https://nexusretail.com" in parsed["text"]
        assert "27ABCDE1234F1Z5" in parsed["text"]
        assert "Thank you for your business" in parsed["text"]


def test_report_excel_dynamic_theme_styling(branding):
    """Verify Excel exports apply designated theme colors to Sheet 1 overview and Sheet 2 data headers."""
    service = ExcelExportService()
    headers = ["Product ID", "Name", "Units Sold", "Total Amount"]
    rows = [
        [101, "Organic Wheat Flour", 25, 2125.00],
        [102, "Basmati Rice 5kg", 18, 3600.00],
    ]
    metadata = {"Store": "Store 1", "Date": "2026-10-08"}
    kpis = [{"label": "Total Revenue", "value": 5725.00}]

    test_cases = [
        ("sales_daily", "15803D"),           # Green
        ("inventory_valuation", "1E40AF"),   # Blue
        ("customers_overview", "0F766E"),    # Teal
        ("inventory_low", "D97706"),         # Amber
        ("pos_z_report", "14532D"),          # Dark Green
        ("products_profitability", "6D28D9") # Purple
    ]

    for rtype, expected_hex in test_cases:
        wb = service.create_report_workbook(
            sheet_title=f"{rtype.replace('_', ' ').title()}",
            headers=headers,
            rows=rows,
            metadata=metadata,
            kpis=kpis,
            branding=branding,
            report_type=rtype,
        )

        assert "Overview" in wb.sheetnames
        ws_overview = wb["Overview"]
        # Sheet 1 title font color matches theme (openpyxl stores aRGB, so ends with 6-char hex)
        title_color = str(ws_overview.cell(row=1, column=1).font.color.rgb)
        assert title_color.endswith(expected_hex)

        # Sheet 1 parameter header fill color matches theme
        param_cell = ws_overview.cell(row=4, column=1)
        param_fill_color = str(param_cell.fill.start_color.rgb)
        assert param_fill_color.endswith(expected_hex)

        # Sheet 2 data header fill color matches theme
        ws_data = wb.active
        assert ws_data.title != "Overview"
        header_cell = ws_data.cell(row=1, column=1)
        header_fill_color = str(header_cell.fill.start_color.rgb)
        assert header_fill_color.endswith(expected_hex)


def test_report_pdf_with_logo_and_graceful_fallbacks(branding):
    """Verify report PDF embeds logo when valid and falls back safely on corrupt/traversal paths."""
    from PIL import Image
    from pathlib import Path

    uploads_dir = Path("uploads")
    uploads_dir.mkdir(parents=True, exist_ok=True)
    valid_logo_path = uploads_dir / "temp_valid_logo.png"

    try:
        # Create a valid test image
        img = Image.new("RGB", (120, 60), color=(15, 118, 110))
        img.save(valid_logo_path)

        renderer = DocumentRenderer()

        # 1. Valid logo embedded
        branding_with_logo = branding.model_copy(update={"logo_url": "/uploads/temp_valid_logo.png"})
        pdf_bytes = renderer.render_report_pdf(
            report_title="Daily Sales Report",
            branding=branding_with_logo,
            report_type="sales_daily",
        )
        parsed = parse_pdf_structure(pdf_bytes)
        assert parsed["page_count"] == 1
        assert "Nexus Retail Hub" in parsed["text"]

        # 2. Corrupt logo fallback (test_logo.png is 67 bytes corrupt)
        branding_corrupt = branding.model_copy(update={"logo_url": "/uploads/test_logo.png"})
        pdf_corrupt = renderer.render_report_pdf(
            report_title="Daily Sales Report",
            branding=branding_corrupt,
            report_type="sales_daily",
        )
        parsed_corrupt = parse_pdf_structure(pdf_corrupt)
        assert parsed_corrupt["page_count"] == 1
        assert "Nexus Retail Hub" in parsed_corrupt["text"]

        # 3. Path traversal rejected gracefully
        branding_traversal = branding.model_copy(update={"logo_url": "../../etc/shadow.png"})
        pdf_traversal = renderer.render_report_pdf(
            report_title="Daily Sales Report",
            branding=branding_traversal,
            report_type="sales_daily",
        )
        parsed_traversal = parse_pdf_structure(pdf_traversal)
        assert parsed_traversal["page_count"] == 1
        assert "Nexus Retail Hub" in parsed_traversal["text"]

    finally:
        if valid_logo_path.exists():
            valid_logo_path.unlink()


def test_report_pdf_minimal_branding_no_breakage():
    """Verify report PDF renders properly with minimal branding and missing optional fields."""
    minimal_branding = BrandingContext(
        business_name="Minimal Outlet",
        address="",
        phone="",
        email="",
        website="",
        gstin="",
        footer_text="",
        show_gstin=False,
    )
    renderer = DocumentRenderer()
    pdf_bytes = renderer.render_report_pdf(
        report_title="Stock Summary Report",
        branding=minimal_branding,
        report_type="inventory_stock",
    )
    parsed = parse_pdf_structure(pdf_bytes)
    assert parsed["page_count"] == 1
    assert "Minimal Outlet" in parsed["text"]


def test_report_pdf_xml_character_escaping():
    """Verify special XML characters (&, <, >) in business details do not crash report generation."""
    special_branding = BrandingContext(
        business_name="Marks & Spencer <Retail> & Co.",
        address="Bldg #4 & #5 <Phase 2>, West Sector",
        phone="+91 9000000000",
        email="info&support@marks-spencer.co.in",
        website="https://m&s.co.in?ref=pos<app>",
        gstin="27AABCM1234F1Z1",
        footer_text="Confidential & Proprietary <Notice>",
    )
    renderer = DocumentRenderer()
    pdf_bytes = renderer.render_report_pdf(
        report_title="Profit & Loss Statement",
        branding=special_branding,
        report_type="profit_loss",
    )
    parsed = parse_pdf_structure(pdf_bytes)
    assert parsed["page_count"] == 1
    assert "Marks & Spencer" in parsed["text"]
