import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl

from app.services.excel_export_service import ExcelExportService


def test_excel_export_service_create_workbook():
    service = ExcelExportService()
    headers = ["Item ID", "Name", "Qty", "Price", "Date"]
    rows = [
        [1, "Product A", 10, Decimal("199.99"), date(2026, 10, 4)],
        [2, "Product B", 5, 49.50, datetime(2026, 10, 4, 12, 0, 0)],
        [3, "Product C with a very long title", 1, 1000, "2026-10-04"],
    ]

    wb = service.create_workbook(sheet_title="Test Sheet", headers=headers, rows=rows)
    assert isinstance(wb, openpyxl.Workbook)

    ws = wb.active
    assert ws.title == "Test Sheet"
    assert ws.cell(row=1, column=1).value == "Item ID"
    assert ws.cell(row=2, column=2).value == "Product A"
    assert ws.cell(row=2, column=4).value == 199.99
    assert ws.cell(row=3, column=3).value == 5

    # Check header styling
    h_cell = ws.cell(row=1, column=1)
    assert h_cell.font.bold is True

    # Check freeze pane
    assert ws.freeze_panes == "A2"


def test_excel_export_service_export_to_bytes_and_stream():
    service = ExcelExportService()
    headers = ["Col A", "Col B"]
    rows = [["Val 1", 100], ["Val 2", 200]]

    # Bytes
    raw_bytes = service.export_to_bytes("Data", headers, rows)
    assert isinstance(raw_bytes, bytes)
    assert len(raw_bytes) > 0
    # Must be valid zip/xlsx signature (starts with PK\x03\x04)
    assert raw_bytes[:4] == b"PK\x03\x04"

    # Stream
    stream = service.export_to_stream("Data", headers, rows)
    assert isinstance(stream, io.BytesIO)
    stream_bytes = stream.getvalue()
    assert stream_bytes[:4] == b"PK\x03\x04"

    # Verify openpyxl can read it back
    parsed_wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    assert parsed_wb.active.cell(row=1, column=1).value == "Col A"
    assert parsed_wb.active.cell(row=2, column=2).value == 100


def test_excel_export_service_streaming_response():
    service = ExcelExportService()
    headers = ["Header 1"]
    rows = [["Data 1"]]

    resp = service.create_streaming_response(
        sheet_title="Stream Test",
        headers=headers,
        rows=rows,
        filename="test_export",
    )
    assert resp.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.headers["Content-Disposition"] == "attachment; filename=test_export.xlsx"
