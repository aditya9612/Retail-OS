import io
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


@pytest.fixture
def export_test_setup(unique_slug):
    email = f"exp-{unique_slug}@test.com"
    client.post(
        "/api/v1/auth/register",
        params={
            "tenant_name": "Export Tenant",
            "slug": unique_slug,
            "email": email,
            "admin_name": "Export Admin",
            "password": "Password123!",
        },
    )
    login = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "Password123!"},
    )
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Store
    store = client.post(
        "/api/v1/stores",
        json={"name": "Export Store", "code": f"E{unique_slug[:3]}"},
        headers=headers,
    ).json()

    # Product
    product = client.post(
        "/api/v1/products",
        json={
            "name": "Export Item",
            "sku": f"SKU-{unique_slug[:4]}",
            "price": "200.00",
            "gst_rate": "18.00",
            "hsn_code": "6109",
        },
        headers=headers,
    ).json()

    # Supplier
    supplier_resp = client.post(
        "/api/v1/suppliers",
        json={
            "name": "Main Supplier Ltd",
            "contact_person": "Vendor Bob",
            "phone": "9876543210",
            "email": f"vendor-{unique_slug[:6]}@supplier.com",
            "address": "45 Industrial Area",
        },
        headers=headers,
    )
    assert supplier_resp.status_code == 201, supplier_resp.text
    supplier = supplier_resp.json()

    # Stock
    client.post(
        "/api/v1/inventory/stock-in",
        json={"store_id": store["id"], "product_id": product["id"], "quantity": 50},
        headers=headers,
    )

    # Cart & Invoice
    client.post(
        "/api/v1/billing/cart/add-item",
        params={"store_id": store["id"], "same_state": True},
        json={"product_id": product["id"], "quantity": "2"},
        headers=headers,
    )
    invoice_resp = client.post(
        "/api/v1/invoices",
        json={
            "store_id": store["id"],
            "same_state": True,
            "payments": [{"payment_mode": "cash", "amount": "472.00"}],
        },
        headers=headers,
    )
    assert invoice_resp.status_code == 201, invoice_resp.text
    invoice = invoice_resp.json()

    # Purchase Order
    po_resp = client.post(
        "/api/v1/purchase-orders",
        json={
            "supplier_id": supplier["id"],
            "store_id": store["id"],
            "items": [
                {
                    "product_id": product["id"],
                    "quantity": 25,
                    "unit_cost": "120.00",
                }
            ],
            "notes": "Urgent procurement",
        },
        headers=headers,
    )
    assert po_resp.status_code == 201, po_resp.text
    po = po_resp.json()

    return {
        "headers": headers,
        "store_id": store["id"],
        "product_id": product["id"],
        "invoice_id": invoice["id"],
        "invoice_number": invoice["invoice_number"],
        "po_id": po["id"],
        "po_number": po["order_number"],
        "slug": unique_slug,
    }


def test_invoice_pdf_tax_invoice_mode(export_test_setup):
    headers = export_test_setup["headers"]
    invoice_id = export_test_setup["invoice_id"]

    res = client.get(
        f"/api/v1/invoices/{invoice_id}/pdf?document_type=invoice&mode=download",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")
    assert b"%%EOF" in res.content
    assert len(res.content) > 1000


def test_invoice_pdf_bill_mode(export_test_setup):
    headers = export_test_setup["headers"]
    invoice_id = export_test_setup["invoice_id"]

    res = client.get(
        f"/api/v1/invoices/{invoice_id}/pdf?document_type=bill&mode=download",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "bill_" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")
    assert b"%%EOF" in res.content
    assert len(res.content) > 1000


def test_purchase_order_pdf(export_test_setup):
    headers = export_test_setup["headers"]
    po_id = export_test_setup["po_id"]
    po_number = export_test_setup["po_number"]

    res = client.get(
        f"/api/v1/purchase-orders/{po_id}/pdf?mode=download",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert f"PO-{po_number}.pdf" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-")
    assert b"%%EOF" in res.content
    assert len(res.content) > 1000


def test_credit_note_pdf(export_test_setup):
    headers = export_test_setup["headers"]
    invoice_id = export_test_setup["invoice_id"]

    # Issue a credit note for this invoice
    cn_res = client.post(
        "/api/v1/credit-notes",
        json={
            "invoice_id": invoice_id,
            "refund_amount": "100.00",
            "reason": "Customer damaged item return",
        },
        headers=headers,
    )
    assert cn_res.status_code == 201
    cn = cn_res.json()

    # Download credit note PDF
    pdf_res = client.get(
        f"/api/v1/credit-notes/{cn['id']}/pdf",
        headers=headers,
    )
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert "CN-" in pdf_res.headers["content-disposition"]
    assert pdf_res.content.startswith(b"%PDF-")
    res_pdf = pdf_res.content
    assert b"%%EOF" in res_pdf
    assert len(res_pdf) > 1000


def test_customer_directory_export_excel_and_pdf(export_test_setup):
    headers = export_test_setup["headers"]

    # 1. Excel export
    res_excel = client.get(
        "/api/v1/customers/export-directory?status=all&format=excel",
        headers=headers,
    )
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.content[:4] == b"PK\x03\x04"

    # Verify openpyxl can load the exported directory
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert "Customer Directory" in wb.sheetnames
    assert wb.active.cell(row=1, column=1).value == "ID"

    # 2. PDF export
    res_pdf = client.get(
        "/api/v1/customers/export-directory?status=all&format=pdf",
        headers=headers,
    )
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")
    assert b"%%EOF" in res_pdf.content


def test_reports_export_excel_and_pdf(export_test_setup):
    headers = export_test_setup["headers"]

    # 1. Reports Excel export
    res_excel = client.post(
        "/api/v1/reports/export",
        json={
            "report_type": "sales_daily",
            "format": "excel",
            "target_date": str(date.today()),
        },
        headers=headers,
    )
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.content[:4] == b"PK\x03\x04"

    # Load with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert wb.active.cell(row=1, column=1).value == "Metric"

    # 2. Reports PDF export
    res_pdf = client.post(
        "/api/v1/reports/export",
        json={
            "report_type": "sales_daily",
            "format": "pdf",
            "target_date": str(date.today()),
        },
        headers=headers,
    )
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")
    assert b"%%EOF" in res_pdf.content


def test_delivery_export_excel_and_csv(export_test_setup):
    headers = export_test_setup["headers"]

    # 1. Delivery Excel export
    res_excel = client.get(
        "/api/v1/delivery/export?format=excel",
        headers=headers,
    )
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.content[:4] == b"PK\x03\x04"

    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert "Deliveries" in wb.sheetnames
    assert wb.active.cell(row=1, column=1).value == "ID"

    # 2. Delivery CSV export
    res_csv = client.get(
        "/api/v1/delivery/export?format=csv",
        headers=headers,
    )
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    assert b"ID,Order Number,Status" in res_csv.content


def test_tenant_isolation_on_document_pdf_exports(export_test_setup, unique_slug):
    # Setup Tenant B
    slug_b = f"b-exp-{unique_slug}"
    email_b = f"exp-b-{unique_slug}@test.com"
    client.post(
        "/api/v1/auth/register",
        params={
            "tenant_name": "Attacker Tenant",
            "slug": slug_b,
            "email": email_b,
            "admin_name": "Attacker",
            "password": "Password123!",
        },
    )
    login_b = client.post(
        "/api/v1/auth/login",
        json={"email": email_b, "password": "Password123!"},
    )
    headers_b = {"Authorization": f"Bearer {login_b.json()['access_token']}"}

    # Tenant B tries to access Tenant A's invoice PDF -> 404
    invoice_id = export_test_setup["invoice_id"]
    res_inv = client.get(f"/api/v1/invoices/{invoice_id}/pdf", headers=headers_b)
    assert res_inv.status_code == 404

    # Tenant B tries to access Tenant A's PO PDF -> 404
    po_id = export_test_setup["po_id"]
    res_po = client.get(f"/api/v1/purchase-orders/{po_id}/pdf", headers=headers_b)
    assert res_po.status_code == 404


def test_reports_get_export_direct_download(export_test_setup):
    headers = export_test_setup["headers"]

    # 1. GET PDF direct download
    res_pdf = client.get(
        "/api/v1/reports/export?report_type=sales_daily&format=pdf",
        headers=headers,
    )
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")
    assert b"%%EOF" in res_pdf.content

    # 2. GET Excel direct download
    res_excel = client.get(
        "/api/v1/reports/export?report_type=sales_daily&format=excel",
        headers=headers,
    )
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.content[:4] == b"PK\x03\x04"

    # 3. GET CSV direct download
    res_csv = client.get(
        "/api/v1/reports/export?report_type=sales_daily&format=csv",
        headers=headers,
    )
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]


def test_expanded_report_types_export(export_test_setup):
    headers = export_test_setup["headers"]

    report_types_to_test = [
        "products_slow_moving",
        "customers_lifetime_value",
        "daily_billing",
        "payment_summary",
        "gst_sales",
        "profit_loss",
        "inventory_stock",
        "inventory_valuation",
    ]

    for rtype in report_types_to_test:
        # Test PDF export
        res_pdf = client.post(
            "/api/v1/reports/export",
            json={"report_type": rtype, "format": "pdf"},
            headers=headers,
        )
        assert res_pdf.status_code == 200, f"PDF failed for {rtype}: {res_pdf.text}"
        assert res_pdf.headers["content-type"] == "application/pdf"
        assert res_pdf.content.startswith(b"%PDF-")

        # Test Excel export
        res_excel = client.post(
            "/api/v1/reports/export",
            json={"report_type": rtype, "format": "excel"},
            headers=headers,
        )
        assert res_excel.status_code == 200, f"Excel failed for {rtype}: {res_excel.text}"
        assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert res_excel.content[:4] == b"PK\x03\x04"


def test_pos_z_report_pdf_and_excel_exports(export_test_setup):
    headers = export_test_setup["headers"]

    from app.core.database import SessionLocal
    from app.models.user import User

    db = SessionLocal()
    try:
        user_db = db.query(User).filter(User.email == f"exp-{export_test_setup['slug']}@test.com").first()
        if user_db:
            user_db.store_id = export_test_setup["store_id"]
            db.commit()
    finally:
        db.close()

    # Open a POS shift
    open_resp = client.post(
        "/api/v1/pos/shifts/open",
        json={"opening_cash_float": "1000.00", "notes": "Shift 1"},
        headers=headers,
    )
    assert open_resp.status_code in (200, 201), open_resp.text
    shift_id = open_resp.json()["id"]

    # 1. Export Z-report PDF
    res_pdf = client.get(
        f"/api/v1/pos/shifts/{shift_id}/z-report/pdf",
        headers=headers,
    )
    assert res_pdf.status_code == 200, res_pdf.text
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")
    assert b"%%EOF" in res_pdf.content

    # 2. Export Z-report Excel
    res_excel = client.get(
        f"/api/v1/pos/shifts/{shift_id}/z-report/excel",
        headers=headers,
    )
    assert res_excel.status_code == 200, res_excel.text
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.content[:4] == b"PK\x03\x04"
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert "Overview" in wb.sheetnames


def test_amount_to_indian_words_utility():
    from app.utils.helpers import amount_to_indian_words

    assert amount_to_indian_words(0) == "Indian Rupees Zero Only"
    assert amount_to_indian_words(1) == "Indian Rupees One Only"
    assert amount_to_indian_words(100) == "Indian Rupees One Hundred Only"
    assert amount_to_indian_words("354.00") == "Indian Rupees Three Hundred Fifty-Four Only"
    assert amount_to_indian_words(1234.50) == "Indian Rupees One Thousand Two Hundred Thirty-Four and Fifty Paise Only"
    assert "Lakh" in amount_to_indian_words(100000)
    assert "Crore" in amount_to_indian_words(10000000)
